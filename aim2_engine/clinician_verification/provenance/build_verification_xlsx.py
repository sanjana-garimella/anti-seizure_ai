#!/usr/bin/env python3
"""Build the clinician verification workbook from the draft data."""
import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HERE=os.path.dirname(os.path.abspath(__file__)); OUT=os.path.dirname(HERE)
draft=json.load(open(os.path.join(OUT,"asm_rules.draft.json")))
rows=json.load(open(os.path.join(OUT,"verification_rows.json")))

NAVY=PatternFill("solid",fgColor="1F3864"); LBLUE=PatternFill("solid",fgColor="D6E4F0")
AMBER=PatternFill("solid",fgColor="FFF2CC"); GREY=PatternFill("solid",fgColor="F2F2F2")
WHITEB=Font(color="FFFFFF",bold=True,size=11); BOLD=Font(bold=True); ITAL=Font(italic=True,color="555555")
WRAP=Alignment(wrap_text=True,vertical="top"); TOP=Alignment(vertical="top")
thin=Side(style="thin",color="BBBBBB"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=openpyxl.Workbook()

# ---------------- READ ME ----------------
ws=wb.active; ws.title="READ ME"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=2; ws.column_dimensions["B"].width=110
def line(r,txt,font=None,fill=None):
    c=ws.cell(row=r,column=2,value=txt); c.alignment=WRAP
    if font:c.font=font
    if fill:c.fill=fill
ws.cell(row=1,column=2,value="Aim 2 - Draft ASM Titration Data: Clinician Verification Workbook").font=Font(bold=True,size=15,color="1F3864")
r=3
for t,f in [
 ("UNVERIFIED DRAFT - FOR DR. BARBOUR'S REVIEW. Nothing here is authoritative.",Font(bold=True,color="C00000")),
 ("Every clinical value in the 'Verification' tab is a DRAFT compiled from public references to give you something to edit rather than a blank page. Please confirm or correct each row before any value is used by the engine.",None),
 ("",None),
 ("How to use this workbook",BOLD),
 ("1. 'Verification' tab: one row per editable clinical value. Set 'Confirm correct? (Y/N)'. If N, put the right value in 'Corrected value' and add a note.",None),
 ("2. 'Formulations' tab: tablet/solution strengths pulled from RxNorm (US NLM) - confirm which are stocked / appropriate.",None),
 ("3. 'Safety & interactions' tab: the safety framing that should drive the engine's safety-check layer - confirm wording.",None),
 ("4. When done, the engineering side transfers your confirmed values into data/asm_rules.json and records sign-off in NEEDS_PI_SIGNOFF.md.",None),
 ("",None),
 ("Provenance",BOLD),
 ("Formulations: RxNorm SCD concepts (rxnav.nlm.nih.gov), pulled 2026-06-17.",None),
 ("Dose bands: FDA prescribing information / standard pediatric dosing references (see 'Source' column per row), pulled 2026-06-17.",None),
 ("Note: openFDA label dataset updates were reported suspended in 2025; values were cross-checked against current FDA PI pages where possible but still require your verification.",ITAL),
 ("",None),
 ("Key things to look at first (flagged in the data)",BOLD),
 ("- Lamotrigine has THREE titration regimens (no valproate / with valproate / with enzyme inducer). The with-valproate regimen is ~50% lower and slower (SJS rash risk). Confirm all three.",None),
 ("- Zonisamide pediatric use (<16y) is OFF-LABEL (FDA-approved >=16y). Confirm you want it in scope and confirm the off-label dosing.",None),
 ("- Ethosuximide covers ABSENCE only (no focal/GTC efficacy).",None),
 ("- Weight-based vs fixed-mg: most drugs are dosed mg/kg/day; clobazam is by weight band (<=30 / >30 kg); carbamazepine/ethosuximide are partly fixed-mg.",None),
 ("- 'absolute_max_mg_per_day' is the ceiling regardless of weight - confirm these caps.",None),
]:
    line(r,t,f); r+=1

# ---------------- VERIFICATION ----------------
vs=wb.create_sheet("Verification"); vs.sheet_view.showGridLines=False
hdr=["Drug","Band / regimen","Field","Draft value (UNVERIFIED)","Source","Confirm correct? (Y/N)","Corrected value","Clinician notes"]
widths=[24,40,34,26,46,16,20,40]
for i,(h,w) in enumerate(zip(hdr,widths),start=1):
    c=vs.cell(row=1,column=i,value=h); c.fill=NAVY; c.font=WHITEB; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=BORD
    vs.column_dimensions[get_column_letter(i)].width=w
vs.freeze_panes="A2"; vs.auto_filter.ref=f"A1:H{1+len(rows)}"
dv=DataValidation(type="list",formula1='"Y,N"',allow_blank=True); vs.add_data_validation(dv)
prev=None
for ridx,row in enumerate(rows,start=2):
    vals=[row["drug"],row["band"],row["field"],str(row["draft_value"]),row["source"],"",row["corrected_value"],row["clinician_notes"]]
    band_change = row["drug"]!=prev
    for i,v in enumerate(vals,start=1):
        c=vs.cell(row=ridx,column=i,value=v); c.alignment=WRAP; c.border=BORD
        if i in (1,) and v: c.font=BOLD
        if band_change: c.fill=LBLUE if i<=5 else GREY
        elif i==6: c.fill=AMBER
        elif i in (7,8): c.fill=GREY
    dv.add(vs.cell(row=ridx,column=6)); prev=row["drug"]
    vs.row_dimensions[ridx].height=30

# ---------------- FORMULATIONS ----------------
fs=wb.create_sheet("Formulations"); fs.sheet_view.showGridLines=False
for i,(h,w) in enumerate(zip(["Drug","RxCUI","Form","Strengths / concentration","In study formulary? (Y/N)","Notes"],[24,10,46,34,20,34]),start=1):
    c=fs.cell(row=1,column=i,value=h); c.fill=NAVY; c.font=WHITEB; c.border=BORD
    fs.column_dimensions[get_column_letter(i)].width=w
fs.freeze_panes="A2"
dv2=DataValidation(type="list",formula1='"Y,N"',allow_blank=True); fs.add_data_validation(dv2)
fr=2
for dkey in [k for k in draft if not k.startswith("_")]:
    d=draft[dkey]; first=True
    for form in d["formulations_rxnorm"]:
        strengths = (", ".join(f"{s} mg" for s in form["strengths_mg"]) if "strengths_mg" in form
                     else f"{form.get('mg_per_ml')} mg/mL")
        vals=[d["display_name"] if first else "", d["rxcui"] if first else "", form["form"], strengths, "", ""]
        for i,v in enumerate(vals,start=1):
            c=fs.cell(row=fr,column=i,value=v); c.alignment=WRAP; c.border=BORD
            if i==1 and v: c.font=BOLD
            if i==5: c.fill=AMBER
        dv2.add(fs.cell(row=fr,column=5)); fr+=1; first=False

# ---------------- SAFETY ----------------
ss=wb.create_sheet("Safety & interactions"); ss.sheet_view.showGridLines=False
for i,(h,w) in enumerate(zip(["Drug","Key safety (DRAFT - confirm)","Key interactions (DRAFT)","Confirm? (Y/N)","Clinician edits"],[24,60,46,14,40]),start=1):
    c=ss.cell(row=1,column=i,value=h); c.fill=NAVY; c.font=WHITEB; c.border=BORD
    ss.column_dimensions[get_column_letter(i)].width=w
ss.freeze_panes="A2"
dv3=DataValidation(type="list",formula1='"Y,N"',allow_blank=True); ss.add_data_validation(dv3)
sr=2
for dkey in [k for k in draft if not k.startswith("_")]:
    d=draft[dkey]
    safety=d.get("SAFETY_PRIORITY",d.get("key_safety",""))
    vals=[d["display_name"],safety,d.get("key_interactions",""),"",""]
    for i,v in enumerate(vals,start=1):
        c=ss.cell(row=sr,column=i,value=v); c.alignment=WRAP; c.border=BORD
        if i==1: c.font=BOLD
        if i==2 and ("BOXED" in str(v) or "OFF-LABEL" in str(v)): c.fill=PatternFill("solid",fgColor="FFD9D9")
        if i==4: c.fill=AMBER
    dv3.add(ss.cell(row=sr,column=4)); ss.row_dimensions[sr].height=60; sr+=1

path=os.path.join(OUT,"ASM_data_clinician_verification.xlsx")
wb.save(path); print("Saved",path,"| verification rows:",len(rows),"| sheets:",wb.sheetnames)
