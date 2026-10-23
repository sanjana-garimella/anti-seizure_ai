"""
build_review_xlsx.py  (2026-07-14)
Flatten review_cases_v2.json into a reviewer-friendly .xlsx for Dr. Barbour.
One row per case: patient-reported inputs -> engine recommendation / AI summary /
draft message / safety flags -> empty (yellow) columns for Kristen's review.

Engine outputs are RULE-BASED and DRAFT (verified:false); this script only
reformats them. In-folder only (confidentiality rule #1).
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "aim2_engine"))
from engine_core import MedicationContext  # noqa: E402

SRC = os.path.join(HERE, sys.argv[1] if len(sys.argv) > 1 else "review_cases_v2.json")
OUT = os.path.join(HERE, sys.argv[2] if len(sys.argv) > 2 else "Kristen_Review_Cases_v2.xlsx")

# ---- humanizers ----
ADHERENCE = {"almost_never": "Almost never misses", "sometimes": "Sometimes misses",
             "stopped": "Stopped the medicine"}
GOAL = {"still_increasing": "Still increasing (not at goal)", "at_goal": "At goal",
        "not_sure": "Not sure"}
IMPACT = {"no_impact": "none", "small": "small/manageable", "medium_big": "difficult but tolerable",
          "big_wants_change": "big — wants change", "medium_tolerable": "medium (tolerable)",
          "medium_alot": "medium but significant", "big": "big impact", "not_a_fit": "not a fit",
          "not_sure_weird_token": "answer not interpretable"}
SE = {"sleepiness": "sleepiness", "irritability": "irritability/mood", "behavior": "behavior changes",
      "acting_confused": "confusion/‘off’", "dizziness": "dizziness", "rash": "rash"}
RASH = {"small": "small rash", "large": "LARGE rash", "blisters": "BLISTERING rash"}
DUR = {"under_1min": "<1 min", "1_3min": "1–3 min", "over_3min": ">3 min",
       "over_5min": ">5 min", "not_sure": "duration not sure"}
CATLABEL = {"no_change": "No change / continue", "escalate": "Increase dose faster",
            "slower_uptitration": "Slow the dose increases", "add_medication": "Add a second medicine",
            "switch_medication": "Switch medicine", "hold": "Hold current dose",
            "lower_dose": "Lower the dose", "insufficient_data": "Need more info (defer)"}
FOCUS = {"titration": "Titration (goal/no-goal)", "adherence": "Adherence", "safety_ceiling": "Dose ceiling",
         "defer": "Insufficient data", "critical_se": "Critical side effect", "switch": "Switch"}


def dose_str(p):
    amt, unit, freq = p.get("dose_amount"), p.get("dose_unit"), p.get("frequency")
    base = f"{amt:g} {unit} × {freq}/day"
    try:
        mg = MedicationContext(medication=p["medication"], dose_amount=amt, dose_unit=unit,
                               frequency_per_day=freq, concentration=p.get("concentration")).total_daily_dose
        if mg is not None:
            base += f"  (= {mg:g} mg/day)"
    except Exception:
        pass
    return base


def seizure_str(ans):
    if str(ans.get("Q3", "")).lower() == "no":
        return "Seizure-free"
    eps = ans.get("seizure_episodes") or []
    parts = []
    for e in eps:
        look = ", ".join(e.get("look") or []) or "seizure"
        bits = [look]
        if e.get("duration"):
            bits.append(DUR.get(e["duration"], e["duration"]))
        if e.get("frequency_bucket"):
            bits.append(e["frequency_bucket"])
        if e.get("count_4wk"):
            bits.append(f"~{e['count_4wk']}/4wk")
        parts.append(" · ".join(bits))
    if not parts and ans.get("Q5"):
        parts.append(f"~{ans['Q5']} {str(ans.get('Q5_unit','')).replace('_',' ')}")
    return "Breakthrough: " + "; ".join(parts) if parts else "Breakthrough (details not given)"


def se_str(ans):
    lst = ans.get("Q9") or []
    names = [SE.get(x, x) for x in lst]
    if ans.get("Q9a"):
        names = [RASH.get(ans["Q9a"], ans["Q9a"]) if n == "rash" else n for n in names]
    impact = ans.get("Q10b_impact") or ans.get("Q10c_impact")
    extras = []
    if ans.get("Q10d_dose_limited_by_se"):
        extras.append("blocks further increase")
    if ans.get("Q10e_min_goal_unreachable_se"):
        extras.append("can't reach min effective dose")
    out = ", ".join(names) if names else "none reported"
    if impact:
        out += f" (impact: {IMPACT.get(impact, impact)})"
    if extras:
        out += " — " + "; ".join(extras)
    return out


def safety_str(c):
    e = c["engine_output"]
    tags = []
    if e.get("safety_concern_flag"):
        tags.append("⚠ SAFETY")
    if e.get("dose_exceeds_max"):
        tags.append("⚠ DOSE OVER MAX")
    for f in e.get("urgent_flags") or []:
        tags.append(f"urgent: {f}")
    for rf in c.get("redflag_categories") or []:
        tags.append(f"free-text flag: {rf}")
    return "\n".join(tags) or "—"


# ---- columns: (header, width, wrap, kristen?) ----
COLS = [
    ("Case", 6, False, False), ("Focus", 20, True, False), ("Scenario (input intent)", 34, True, False),
    ("Priority", 9, False, False), ("Age", 5, False, False), ("Wt (kg)", 7, False, False),
    ("Medication", 15, False, False), ("Dose (reported)", 22, True, False),
    ("Adherence", 18, True, False), ("Seizures", 26, True, False), ("Side effects", 30, True, False),
    ("Goal status", 20, True, False), ("Can ↑?", 7, False, False),
    ("Recommendation (DRAFT)", 22, True, False), ("Category", 16, False, False),
    ("Engine reason", 40, True, False), ("Safety flags", 22, True, False),
    ("Follow-up / labs / pace", 22, True, False), ("Switch target", 14, True, False),
    ("AI summary (DRAFT — rules-based)", 74, True, False),
    ("Draft family message (DRAFT)", 60, True, False),
    ("Kristen: agree? (Y/N)", 12, True, True), ("Kristen: corrected recommendation", 24, True, True),
    ("Kristen: safety flag OK? (Y/N)", 13, True, True), ("Kristen: comments / clinical notes", 34, True, True),
    ("Kristen: reviewed by / date", 18, True, True),
]

PRIO_FILL = {"urgent": "F4CCCC", "review": "FFF2CC", "routine": "D9EAD3"}
HDR_FILL = PatternFill("solid", fgColor="1F3864")
KHDR_FILL = PatternFill("solid", fgColor="7F6000")
KCELL_FILL = PatternFill("solid", fgColor="FFF9E0")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def followup_str(e):
    bits = []
    if e.get("followup_weeks"):
        bits.append(f"{e['followup_weeks']} wks")
    if e.get("lab_orders"):
        bits.append("labs: " + ", ".join(e["lab_orders"]))
    if e.get("pace") and e.get("pace") != "standard":
        bits.append(f"pace: {e['pace']}")
    if e.get("dosing_note"):
        bits.append(e["dosing_note"])
    return "\n".join(bits) or "—"


def main():
    data = json.load(open(SRC))
    cases = data["cases"]

    wb = Workbook()

    # ---- Read-me / legend sheet ----
    rm = wb.active
    rm.title = "Read me"
    rm.sheet_view.showGridLines = False
    legend = [
        ("Aim 2 — synthetic case review for Dr. Barbour", True),
        ("", False),
        (f"{data['n_cases']} synthetic cases · generated {data['generated']} · engine v{data.get('engine_version')}", False),
        ("Focus: dose-titration goal/no-goal adjustments + critical side effects.", False),
        ("", False),
        ("100% SYNTHETIC — no PHI. The recommendation, AI summary, draft family message, and every", False),
        ("safety flag are produced by the RULE-BASED engine (engine_core) via the production path", False),
        ("(build_snapshot → run_engine → red-flag scan → triage). Nothing clinical is invented here.", False),
        ("Every output is a DRAFT (verified:false), pending PI sign-off (NEEDS_PI_SIGNOFF.md).", False),
        ("", False),
        ("How to review: read each row's inputs + AI summary, then fill the amber columns on the right:", True),
        ("   • Kristen: agree? (Y/N) — Y if the engine's recommendation is clinically correct", False),
        ("   • Kristen: corrected recommendation — if N, what it should be (e.g. 'switch', 'hold')", False),
        ("   • Kristen: safety flag OK? (Y/N) — did it flag/relay urgent items correctly?", False),
        ("   • Kristen: comments — any note for the engineering team", False),
        ("   • Kristen: reviewed by / date — your initials + date", False),
        ("", False),
        ("Example filled row:  Y  |  (blank)  |  Y  |  'agree; would also order a drug level'  |  KB 2026-07-15", False),
        ("", False),
        (f"Category spread: {json.dumps(data['category_counts'])}", False),
        (f"Priority spread: {json.dumps(data['priority_counts'])}", False),
    ]
    for i, (text, bold) in enumerate(legend, start=1):
        c = rm.cell(row=i, column=1, value=text)
        c.font = Font(name="Arial", size=13 if (i == 1) else 10, bold=bold,
                      color="1F3864" if i == 1 else "000000")
    rm.column_dimensions["A"].width = 105

    # ---- Cases sheet ----
    ws = wb.create_sheet("Cases")
    ws.sheet_view.showGridLines = False
    for j, (hdr, width, wrap, kris) in enumerate(COLS, start=1):
        col = get_column_letter(j)
        ws.column_dimensions[col].width = width
        cell = ws.cell(row=1, column=j, value=hdr)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = KHDR_FILL if kris else HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"

    r = 2
    for c in cases:
        p, ans, e = c["patient"], c["answers"], c["engine_output"]
        row = [
            c["case_id"], FOCUS.get(c["group"], c["group"]), c["scenario"], c["priority"].upper(),
            p.get("age"), p.get("weight"), p.get("medication", "").replace("_", " ").title(),
            dose_str(p), ADHERENCE.get(ans.get("Q2"), ans.get("Q2", "—")),
            seizure_str(ans), se_str(ans), GOAL.get(ans.get("Q10a"), ans.get("Q10a", "—")),
            "Yes" if p.get("can_increase") else "No",
            CATLABEL.get(e["category"], e["category"]), e["category"],
            e.get("reason_summary") or "—", safety_str(c), followup_str(e),
            e.get("switch_target") or "—", e.get("clinical_summary") or "",
            e.get("draft_patient_message") or "",
            "", "", "", "", "",  # Kristen input
        ]
        for j, val in enumerate(row, start=1):
            hdr, width, wrap, kris = COLS[j - 1]
            cell = ws.cell(row=r, column=j, value=val)
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(wrap_text=wrap, vertical="top",
                                       horizontal="center" if hdr in ("Priority", "Age", "Wt (kg)", "Can ↑?") else "left")
            cell.border = BORDER
            if kris:
                cell.fill = KCELL_FILL
        # priority color
        pc = ws.cell(row=r, column=4)
        pc.fill = PatternFill("solid", fgColor=PRIO_FILL.get(c["priority"], "FFFFFF"))
        pc.font = Font(name="Arial", size=9, bold=True)
        # row height ~ fit the longer of AI summary / message
        summ = (e.get("clinical_summary") or "")
        msg = (e.get("draft_patient_message") or "")
        lines = max(sum(max(1, (len(ln) // 88) + 1) for ln in summ.split("\n")),
                    sum(max(1, (len(ln) // 72) + 1) for ln in msg.split("\n")))
        ws.row_dimensions[r].height = min(430, max(46, lines * 12.2))
        r += 1

    wb.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
